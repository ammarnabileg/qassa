"""Read an uploaded .xlsx or .csv sheet into rows keyed by canonical column names (Arabic or English headers)."""
import csv
import io

from apps.core.errors import DomainError

MAX_ROWS = 5000


def _norm(h) -> str:
    return " ".join(str(h or "").strip().lower().replace("_", " ").split())


def read_rows(upload, aliases: dict[str, set[str]]) -> list[dict]:
    """aliases: canonical key -> accepted header spellings. Returns [{key: str value}] with the sheet's row number in '_row'."""
    name = (upload.name or "").lower()
    raw = upload.read()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        except Exception:
            raise DomainError(code="import_bad_file")
        table = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        for enc in ("utf-8-sig", "cp1256"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise DomainError(code="import_bad_file")
        table = list(csv.reader(io.StringIO(text)))
    else:
        raise DomainError(code="import_bad_file")
    while table and not any(c not in (None, "") for c in table[0]):
        table.pop(0)
    if not table:
        raise DomainError(code="import_empty")
    lookup = {_norm(a): key for key, names in aliases.items() for a in names | {key}}
    cols = {i: lookup[_norm(h)] for i, h in enumerate(table[0]) if _norm(h) in lookup}
    if not cols:
        raise DomainError(code="import_no_columns", details={"expected": sorted(aliases)})
    rows = []
    for n, r in enumerate(table[1:], start=2):
        row = {k: ("" if r[i] is None else (r[i].isoformat()[:10] if hasattr(r[i], "isoformat") else str(r[i]).strip()))
               for i, k in cols.items() if i < len(r)}
        if any(row.values()):
            row["_row"] = n
            rows.append(row)
    if len(rows) > MAX_ROWS:
        raise DomainError(code="import_too_many", details={"max": MAX_ROWS})
    return rows
